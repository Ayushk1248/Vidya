import os,secrets,io,csv,base64,zipfile,datetime
from fastapi import FastAPI,Request,HTTPException
from fastapi.responses import Response
from openpyxl import load_workbook,Workbook
from .engine import validate
app=FastAPI(docs_url=None,redoc_url=None,openapi_url=None)
FIELDS=['student_id','apaar','name','dob','program_id','admission_date','scheme','category','nationality','mobile','email','passport','gender','year_of_study','admission_mode','social_category','nri_status','oci_status','fn_status','enrolment_status','status_date','apaar_pending_reason','apaar_pending_deadline','parent_name']
@app.middleware('http')
async def guard(request:Request,call_next):
    token=os.environ.get('VALIDATION_SERVICE_TOKEN','')
    if not token or not secrets.compare_digest(request.headers.get('x-service-token',''),token): return Response(status_code=403)
    if int(request.headers.get('content-length','0'))>24_000_000: return Response(status_code=413)
    return await call_next(request)
@app.get('/health')
def health(): return {'ok':True}
def parse(data,filename,academic_year=None):
    if len(data)>8_000_000: raise ValueError('Maximum file size is 8 MB.')
    if filename.lower().endswith('.xlsx'):
        z=zipfile.ZipFile(io.BytesIO(data))
        if sum(i.file_size for i in z.infolist())>30_000_000: raise ValueError('Workbook expands beyond the allowed size.')
        w=load_workbook(io.BytesIO(data),read_only=True,data_only=False); sheet=w.active
        if sheet.max_row and sheet.max_row>10001: raise ValueError('Maximum 10,000 rows.')
        rows=[]
        for cells in sheet.iter_rows():
            if any(c.data_type=='f' for c in cells): raise ValueError('Formulas are not allowed; paste values instead.')
            rows.append([c.value.isoformat()[:10] if isinstance(c.value,(datetime.date,datetime.datetime)) else str(c.value if c.value is not None else '').strip() for c in cells])
            if len(rows)>10001: raise ValueError('Maximum 10,000 rows.')
        w.close()
    elif filename.lower().endswith('.csv'): rows=list(csv.reader(io.StringIO(data.decode('utf-8-sig'))))
    else: raise ValueError('Upload a CSV or XLSX file.')
    if not rows: raise ValueError('File is empty.')
    aliases={'apaar_id':'apaar','student_name':'name','date_of_birth':'dob','program_code':'program_id','seat_type':'scheme','passport_no':'passport'}
    headers=[aliases.get(h.strip().lower(),h.strip().lower()) for h in rows[0]]
    if len(set(headers))!=len(headers):raise ValueError('Duplicate column names after mapping reference aliases. Use each field once.')
    required={'student_id','apaar','name','dob','program_id','admission_date','scheme','nationality'}
    missing=required-set(headers)
    if missing:raise ValueError('Missing required columns: '+', '.join(sorted(missing))+'. Include student_id; academic identities are never invented during import.')
    result=[]
    for values in rows[1:]:
        if not any(values):continue
        if len(values)>len(headers):raise ValueError('A row has extra columns. Quote values containing commas.')
        source=dict(zip(headers,values))
        if academic_year is not None and source.get('academic_year') and source['academic_year'].strip() not in (str(academic_year),f'{academic_year}-{str(int(academic_year)+1)[-2:]}',f'{academic_year}-{int(academic_year)+1}'):
            raise ValueError('The academic_year column must match the selected upload year.')
        r={k:source.get(k,'').strip() for k in FIELDS if k in headers}
        seat=r.get('scheme','');normalized=seat.casefold().replace(' ','').replace('_','-')
        if normalized in ('within','withinintake','within-intake'):r['scheme']='WITHIN'
        elif normalized.startswith('supernumerary-'):r['scheme']=seat.split('-',1)[1].strip().upper()
        else:r['scheme']=seat.upper()
        for field in ('fn_status','oci_status','nri_status'):
            if field in r:r[field]=r[field].upper()
        if not r.get('category'):r['category']='FN' if r.get('fn_status')=='Y' else 'OCI' if r.get('oci_status')=='Y' else 'NRI' if r.get('nri_status')=='Y' else 'GENERAL'
        if r.get('nationality','').casefold() in ('indian','india','in'):r['nationality']='IN'
        result.append(r)
    for r in result:
        defaults={'gender':'U','year_of_study':'1','admission_mode':'Regular','social_category':'GEN','nri_status':'N','oci_status':'Y' if r.get('category')=='OCI' else 'N','fn_status':'Y' if r.get('category')=='FN' else 'N','enrolment_status':'Active','status_date':r.get('admission_date','')}
        for k,v in defaults.items():r[k]=r.get(k) or v
    if not result or len(result)>10000: raise ValueError('Upload between 1 and 10,000 rows.')
    if any(len(v)>250 or v.startswith(('=','+','@')) for r in result for v in r.values()): raise ValueError('Cells must be plain text, at most 250 characters.')
    return result
@app.post('/parse')
async def parse_api(request:Request):
    try:
        body=await request.json(); return {'rows':parse(base64.b64decode(body['data'],validate=True),body['filename'],body.get('academic_year'))}
    except Exception as e: raise HTTPException(400,str(e) if isinstance(e,ValueError) else 'Cannot parse this file. Use the provided template.')
@app.post('/validate')
async def validation(request:Request):
    body=await request.json(); return {'results':validate(body['rows'],body['context'])}
@app.post('/template')
async def template(request:Request):
    body=await request.json(); w=Workbook(); s=w.active; s.title='Synthetic students';s.append(FIELDS)
    for row in body['rows']: s.append([str(row.get(k,'')) for k in FIELDS])
    for col in s.columns:
        for cell in col: cell.number_format='@'
        s.column_dimensions[col[0].column_letter].width=24
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.utils import get_column_letter
    choices={'gender':'M,F,T','year_of_study':'1,2,3,4,5','admission_mode':'Regular,Lateral','social_category':'GEN,EWS,OBC,SC,ST','nri_status':'Y,N','oci_status':'Y,N','fn_status':'Y,N','enrolment_status':'Active,Left,Transferred,Graduated'}
    for field,options in choices.items():
        column=get_column_letter(FIELDS.index(field)+1);dv=DataValidation(type='list',formula1='"'+options+'"',allow_blank=True);s.add_data_validation(dv);dv.add(f'{column}2:{column}10001')
    s.freeze_panes='A2'; b=io.BytesIO();w.save(b)
    return Response(b.getvalue(),media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

# Keep model fitting off the ASGI event loop so upload validation stays responsive.
from starlette.concurrency import run_in_threadpool
@app.post('/anomalies')
async def anomaly_analysis(request:Request):
    payload=await request.json()
    def run():
        from .anomalies import analyse
        return analyse(payload)
    return await run_in_threadpool(run)

@app.post('/export')
async def export_xlsx(request:Request):
    body=await request.json()
    if len(body.get('rows',[]))>50000:raise HTTPException(400,'Maximum 50,000 report rows.')
    w=Workbook();s=w.active;s.title='VIDYA report';s.append(body['columns'])
    for row in body['rows']:s.append([str(v) for v in row])
    from openpyxl.styles import Font,PatternFill
    for cell in s[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='142B49')
    for col in s.columns:s.column_dimensions[col[0].column_letter].width=min(55,max(18,max(len(str(c.value or '')) for c in col[:200])+2))
    s.freeze_panes='A2';s.auto_filter.ref=s.dimensions;b=io.BytesIO();w.save(b)
    return Response(b.getvalue(),media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
